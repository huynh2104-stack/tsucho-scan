"""Test form Excel kiểu Cowork của export.py."""

from datetime import date

from openpyxl import load_workbook

from export import export_accounts

HEADER_ROW = 7


def _tx(out=None, inn=None, bal=None, status="ok", reason="", **kw):
    row = {
        "date": "2025-08-27",
        "date_raw": "07-08-27",
        "kind": "200",
        "description": "テスト",
        "amount_out": out,
        "amount_in": inn,
        "balance": bal,
        "extra_text": None,
        "store_code": None,
        "is_note": False,
        "confidence": "high",
        "status": status,
        "reason": reason,
    }
    row.update(kw)
    return row


def _acc(txs, bank="テスト銀行", **meta):
    acc = {"bank_name": bank, "store_no": None, "account_no": None,
           "holder": None, "transactions": txs}
    acc.update(meta)
    return acc


def _headers(ws):
    return [ws.cell(row=HEADER_ROW, column=c).value
            for c in range(1, ws.max_column + 1)]


def test_khoi_thong_tin_va_cot(tmp_path):
    txs = [
        _tx(out=100, bal=900, extra_text="(282)", store_code="709"),
        _tx(inn=50, bal=950),
    ]
    out = tmp_path / "f.xlsx"
    export_accounts([_acc(txs, store_no="000", account_no="123456",
                          holder="株式会社 見本商事")], str(out), source_name="p1.jpg")
    ws = load_workbook(out).active

    # Khối thông tin
    assert ws["A1"].value == "銀行" and ws["B1"].value == "テスト銀行"
    assert ws["B2"].value == "000"
    assert ws["B3"].value == "123456"
    assert ws["B5"].value == "p1.jpg"

    # Cột động: có cả 備考 và 取扱店 vì dữ liệu có
    hs = _headers(ws)
    assert "備考" in hs and "取扱店" in hs
    assert hs[0] == "年月日(西暦)" and "残高チェック" in hs

    # Ngày là ô date thật
    assert isinstance(ws.cell(row=HEADER_ROW + 1, column=1).value, date)


def test_cot_bien_mat_khi_khong_co_du_lieu(tmp_path):
    out = tmp_path / "f2.xlsx"
    export_accounts([_acc([_tx(out=1, bal=9)])], str(out))
    hs = _headers(load_workbook(out).active)
    assert "備考" not in hs and "取扱店" not in hs


def test_cong_thuc_zandaka_check(tmp_path):
    txs = [_tx(out=100, bal=900), _tx(inn=50, bal=950)]
    out = tmp_path / "f3.xlsx"
    export_accounts([_acc(txs)], str(out))
    ws = load_workbook(out).active
    hs = _headers(ws)
    check_col = hs.index("残高チェック") + 1

    assert ws.cell(row=HEADER_ROW + 1, column=check_col).value == "―"
    formula = ws.cell(row=HEADER_ROW + 2, column=check_col).value
    # Tham chiếu dòng có 残高 phía trên: =IF(E9=E8-N(C9)+N(D9),"OK","NG")
    assert formula.startswith("=IF(") and '"OK"' in formula and '"NG"' in formula


def test_dong_ghi_chu_in_nghieng(tmp_path):
    txs = [
        _tx(out=100, bal=900),
        _tx(is_note=True, status="note", kind=None, date=None, date_raw=None,
            description="新通帳へ繰越しました", bal=None),
    ]
    out = tmp_path / "f4.xlsx"
    export_accounts([_acc(txs)], str(out))
    ws = load_workbook(out).active
    r = HEADER_ROW + 2
    assert ws.cell(row=r, column=1).value == "—"
    assert ws.cell(row=r, column=2).value == "注記"
    assert ws.cell(row=r, column=3).font.italic


def test_bao_cao_va_uchiwake(tmp_path):
    txs = [
        _tx(out=100, bal=900, description="NCB VISA"),
        _tx(inn=50, bal=950, description="ATM入金", kind="100"),
        _tx(inn=70, bal=1020, description="ATM入金", kind="100"),
    ]
    out = tmp_path / "f5.xlsx"
    export_accounts([_acc(txs)], str(out))
    ws = load_workbook(out).active

    vals = {}
    for r in range(HEADER_ROW + len(txs) + 1, ws.max_row + 1):
        label = ws.cell(row=r, column=2).value
        if label:
            vals[label] = (ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value)

    assert vals["期末残高"][0].startswith("=LOOKUP(9.9E+307")
    assert vals["お預り合計"][0].startswith("=SUM(")
    assert "検算（期首＋預り−支払＝期末）" in vals
    # 内訳 sống bằng COUNTIFS/SUMIFS
    assert vals["ATM入金"][0].startswith("=COUNTIFS(")
    assert vals["ATM入金"][1].startswith("=SUMIFS(")
    assert vals["NCB VISA"][0].startswith("=COUNTIFS(")


def test_o_khong_chac_chan_to_vang(tmp_path):
    txs = [_tx(out=100, bal=900, confidence="low"),
           _tx(inn=50, bal=999, status="check", reason="残高不一致(差額 +49)")]
    out = tmp_path / "f6.xlsx"
    export_accounts([_acc(txs)], str(out))
    ws = load_workbook(out).active
    hs = _headers(ws)

    desc_col = hs.index("摘要") + 1
    bal_col = hs.index("差引残高 (円)") + 1
    assert str(ws.cell(row=HEADER_ROW + 1, column=desc_col).fill.start_color.rgb).endswith("FFF2CC")
    assert str(ws.cell(row=HEADER_ROW + 2, column=bal_col).fill.start_color.rgb).endswith("FFF2CC")