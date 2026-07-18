"""Xuất Excel theo form người dùng đã duyệt (học từ skill bank-passbook-to-excel):

- Khối thông tin đầu sheet: 銀行 / 店番号 / 口座番号 / 名義 / 出典 + chú giải màu.
- Cột theo bố cục sổ: 年月日(西暦) | 区分 | 摘要 | (備考) | お支払 | お預り | 差引残高
  | (取扱店) | 残高チェック | 確認メモ — cột 備考/取扱店 chỉ hiện khi sổ có dữ liệu đó.
- 残高チェック: công thức Excel sống từng dòng, tham chiếu dòng có 残高 gần nhất.
- Ô chưa chắc chắn tô FFF2CC; ô được kiểm tra kép tự sửa tô xanh C6EFCE.
- Dòng ghi chú (繰越, 未記帳...) in nghiêng xám, không có số tiền.
- Báo cáo 集計レポート + 入金/出金の内訳 bằng công thức COUNTIFS/SUMIFS sống.
"""

import re
from collections import OrderedDict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

UNCERTAIN = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FIXED_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="2F5B7C", end_color="2F5B7C", fill_type="solid")
REPORT_FILL = PatternFill(start_color="E8EEF4", end_color="E8EEF4", fill_type="solid")
BAND_FILL = PatternFill(start_color="E4EEF6", end_color="E4EEF6", fill_type="solid")  # dòng so le

ARIAL = Font(name="Arial")
ARIAL_BOLD = Font(name="Arial", bold=True)
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
NOTE_FONT = Font(name="Arial", italic=True, color="808080")
FAINT_FONT = Font(name="Arial", color="9AA5AF")  # OK/NG/― viết nhạt

CENTER = Alignment(horizontal="center")
# Tên riêng katakana (≥3 ký tự liền) — chỗ model hay sai chính tả mà không tự biết
_KANA_RUN = re.compile(r"[ァ-ヶー]{3,}")

MONEY_FMT = "#,##0"
DATE_FMT = "yyyy-mm-dd"
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_ROW = 7  # bảng bắt đầu dưới khối thông tin


def _sheet_title(bank_name: str | None, index: int) -> str:
    """Tên sheet hợp lệ cho Excel (bỏ ký tự cấm, tối đa 31 ký tự)."""
    name = bank_name or f"口座{index + 1}"
    name = re.sub(r"[\\/*?:\[\]]", "", name).strip() or f"口座{index + 1}"
    return name[:28]


def _escape_countif(key: str) -> str:
    """Escape cho COUNTIFS/SUMIFS: nhân đôi nháy kép, vô hiệu wildcard * ? ~."""
    s = key.replace("~", "~~").replace("*", "~*").replace("?", "~?")
    return s.replace('"', '""')


def _set(ws, row: int, col: int, value, font=ARIAL, fill=None, fmt=None, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    return cell


def _write_sheet(ws, acc: dict, source_name: str) -> None:
    txs = acc.get("transactions", [])
    has_x = any(t.get("extra_text") for t in txs)
    has_t = any(t.get("store_code") for t in txs)

    # ---- Khối thông tin đầu sheet
    info = [
        ("銀行", acc.get("bank_name")),
        ("店番号", acc.get("store_no")),
        ("口座番号", acc.get("account_no")),
        ("名義", acc.get("holder")),
        ("出典", source_name),
    ]
    for r, (label, val) in enumerate(info, start=1):
        _set(ws, r, 1, label, font=ARIAL_BOLD)
        _set(ws, r, 2, val or "―")
    _set(ws, 1, 6, "Ô vàng nhạt = chưa chắc chắn / tên katakana cần soát（黄色＝要確認）",
         fill=UNCERTAIN)
    _set(ws, 2, 6, "Ô xanh lá = đã tự sửa bằng kiểm tra số dư（緑＝残高検算で自動修正）",
         fill=FIXED_FILL)

    # ---- Cột động theo bố cục sổ
    headers = ["年月日(西暦)", "区分", "摘要"]
    if has_x:
        headers.append("備考")
    headers += ["お支払金額 (円)", "お預り金額 (円)", "差引残高 (円)"]
    if has_t:
        headers.append("取扱店")
    headers += ["残高チェック", "確認メモ"]

    col = {name: i + 1 for i, name in enumerate(headers)}
    letter = {name: get_column_letter(i) for name, i in col.items()}
    out_l, in_l, bal_l = letter["お支払金額 (円)"], letter["お預り金額 (円)"], letter["差引残高 (円)"]

    for name, c in col.items():
        _set(ws, HEADER_ROW, c, name, font=HEADER_FONT, fill=HEADER_FILL, border=BORDER)
        ws.cell(row=HEADER_ROW, column=c).alignment = Alignment(horizontal="center")
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    # ---- Dòng dữ liệu
    prev_bal_row = None
    r = HEADER_ROW
    for tx in txs:
        r += 1
        is_note = tx.get("is_note")
        status = tx.get("status")

        # Ngày: ô date thật (yyyy-mm-dd); không đọc được thì giữ chuỗi gốc + tô vàng
        date_cell = _set(ws, r, col["年月日(西暦)"], None, border=BORDER)
        if is_note:
            date_cell.value = "—"
        elif tx.get("date"):
            date_cell.value = date.fromisoformat(tx["date"])
            date_cell.number_format = DATE_FMT
        else:
            date_cell.value = tx.get("date_raw") or ""
            date_cell.fill = UNCERTAIN

        _set(ws, r, col["区分"], tx.get("kind") or ("注記" if is_note else None), border=BORDER)
        desc_cell = _set(ws, r, col["摘要"], tx.get("description") or None, border=BORDER)
        if has_x:
            _set(ws, r, col["備考"], tx.get("extra_text") or None, border=BORDER)

        out_cell = _set(ws, r, col["お支払金額 (円)"], tx.get("amount_out"),
                        fmt=MONEY_FMT, border=BORDER)
        in_cell = _set(ws, r, col["お預り金額 (円)"], tx.get("amount_in"),
                       fmt=MONEY_FMT, border=BORDER)
        bal_cell = _set(ws, r, col["差引残高 (円)"], tx.get("balance"),
                        fmt=MONEY_FMT, border=BORDER)
        if has_t:
            _set(ws, r, col["取扱店"], tx.get("store_code") or None, border=BORDER)

        # 残高チェック: công thức sống, tham chiếu dòng có 残高 gần nhất phía trên
        if tx.get("balance") is not None and prev_bal_row is not None:
            check = (f'=IF({bal_l}{r}={bal_l}{prev_bal_row}'
                     f'-N({out_l}{r})+N({in_l}{r}),"OK","NG")')
        else:
            check = "―"
        check_cell = _set(ws, r, col["残高チェック"], check, font=FAINT_FONT, border=BORDER)
        check_cell.alignment = CENTER
        if tx.get("balance") is not None:
            prev_bal_row = r

        # Căn giữa các cột mã/ngày cho dễ đọc
        date_cell.alignment = CENTER
        ws.cell(row=r, column=col["区分"]).alignment = CENTER
        if has_t:
            ws.cell(row=r, column=col["取扱店"]).alignment = CENTER
        if has_x:
            ws.cell(row=r, column=col["備考"]).alignment = CENTER

        # Ghi chú xác nhận + tô màu
        if status == "check":
            _set(ws, r, col["確認メモ"], f"要確認: {tx.get('reason', '')}", border=BORDER)
            for cell in (out_cell, in_cell, bal_cell):
                cell.fill = UNCERTAIN
        elif status == "fixed":
            _set(ws, r, col["確認メモ"], f"自動修正: {tx.get('reason', '')}", border=BORDER)
            for cell in (out_cell, in_cell):
                cell.fill = FIXED_FILL
        else:
            _set(ws, r, col["確認メモ"], None, border=BORDER)

        # Tên katakana (model hay sai chính tả mà không tự biết) → tô nhạt để soát
        if tx.get("confidence") == "low" or (
            not is_note and _KANA_RUN.search(str(tx.get("description") or ""))
        ):
            desc_cell.fill = UNCERTAIN

        if is_note:
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).font = NOTE_FONT
        else:
            # Dòng so le màu nền nhẹ cho dễ dò ngang (không đè lên ô đã tô màu)
            if (r - HEADER_ROW) % 2 == 0:
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(row=r, column=c)
                    if cell.fill is None or cell.fill.fill_type is None:
                        cell.fill = BAND_FILL

    first, last = HEADER_ROW + 1, r
    bal_rows = [
        HEADER_ROW + 1 + i for i, t in enumerate(txs) if t.get("balance") is not None
    ]

    # ---- 集計レポート (công thức sống)
    r = last + 2
    _set(ws, r, 2, f"■ 集計レポート（{acc.get('bank_name') or 'この口座'}）", font=ARIAL_BOLD)
    report_rows = {}
    items = []
    if bal_rows:
        f_b = bal_rows[0]
        items.append(("期首残高（初回取引前）",
                      f"={bal_l}{f_b}+N({out_l}{f_b})-N({in_l}{f_b})"))
        items.append(("期末残高", f"=LOOKUP(9.9E+307,{bal_l}{first}:{bal_l}{last})"))
    items += [
        ("お預り件数", f"=COUNT({in_l}{first}:{in_l}{last})"),
        ("お預り合計", f"=SUM({in_l}{first}:{in_l}{last})"),
        ("お支払件数", f"=COUNT({out_l}{first}:{out_l}{last})"),
        ("お支払合計", f"=SUM({out_l}{first}:{out_l}{last})"),
    ]
    for label, formula in items:
        r += 1
        report_rows[label] = r
        _set(ws, r, 2, label, fill=REPORT_FILL)
        _set(ws, r, 3, formula, fill=REPORT_FILL, fmt=MONEY_FMT)
    if bal_rows:
        r += 1
        _set(ws, r, 2, "増減（期末－期首）", fill=REPORT_FILL)
        _set(ws, r, 3, f"=C{report_rows['期末残高']}-C{report_rows['期首残高（初回取引前）']}",
             fill=REPORT_FILL, fmt=MONEY_FMT)
        r += 1
        _set(ws, r, 2, "検算（期首＋預り−支払＝期末）", fill=REPORT_FILL)
        _set(ws, r, 3,
             f"=IF(C{report_rows['期首残高（初回取引前）']}"
             f"+C{report_rows['お預り合計']}-C{report_rows['お支払合計']}"
             f"=C{report_rows['期末残高']},\"一致 ✓\",\"不一致 ✗\")",
             fill=REPORT_FILL)
    else:
        r += 1
        _set(ws, r, 2, "検算", fill=REPORT_FILL)
        _set(ws, r, 3, "残高が読取れず検算不可", fill=REPORT_FILL)

    # ---- 入金・出金の内訳 theo đối tác (COUNTIFS/SUMIFS sống)
    for title, amount_key, amt_l in (
        ("― 入金の内訳（chi tiết tiền vào）", "amount_in", in_l),
        ("― 出金の内訳（chi tiết tiền ra）", "amount_out", out_l),
    ):
        keys: "OrderedDict[tuple, int]" = OrderedDict()
        for t in txs:
            if t.get("is_note") or not t.get(amount_key):
                continue
            name = (t.get("description") or "").strip()
            match_col = "摘要"
            if not name:
                name = (t.get("kind") or "").strip()
                match_col = "区分"
            if not name:
                continue
            keys[(name, match_col)] = keys.get((name, match_col), 0) + t[amount_key]
        if not keys:
            continue

        r += 2
        _set(ws, r, 2, title, font=ARIAL_BOLD)
        r += 1
        for c, label in ((2, "相手先 / 種類"), (3, "件数"), (4, "合計 (円)")):
            _set(ws, r, c, label, font=HEADER_FONT, fill=HEADER_FILL)
        for (name, match_col), _total in sorted(keys.items(), key=lambda kv: -kv[1]):
            r += 1
            m_l = letter[match_col]
            key = _escape_countif(name)
            _set(ws, r, 2, name)
            _set(ws, r, 3,
                 f'=COUNTIFS({m_l}{first}:{m_l}{last},"{key}",'
                 f'{amt_l}{first}:{amt_l}{last},">0")')
            _set(ws, r, 4,
                 f'=SUMIFS({amt_l}{first}:{amt_l}{last},'
                 f'{m_l}{first}:{m_l}{last},"{key}")', fmt=MONEY_FMT)

    # ---- Độ rộng cột
    widths = {"年月日(西暦)": 12, "区分": 14, "摘要": 26, "備考": 10,
              "お支払金額 (円)": 14, "お預り金額 (円)": 14, "差引残高 (円)": 14,
              "取扱店": 8, "残高チェック": 12, "確認メモ": 30}
    for name, c in col.items():
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 12)


def export_accounts(accounts: list[dict], output_path: str, source_name: str = "") -> str:
    """Ghi file .xlsx, mỗi tài khoản (sổ) một sheet riêng theo form đã duyệt."""
    wb = Workbook()
    used_titles = set()
    for i, acc in enumerate(accounts):
        title = _sheet_title(acc.get("bank_name"), i)
        if title in used_titles:
            title = f"{title[:25]}({i + 1})"
        used_titles.add(title)

        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        _write_sheet(ws, acc, source_name)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return str(out)


def export_to_excel(transactions: list[dict], output_path: str) -> str:
    """Giữ tương thích: xuất 1 danh sách giao dịch thành 1 sheet."""
    return export_accounts(
        [{"bank_name": None, "transactions": transactions}], output_path
    )
