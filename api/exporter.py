"""Xuất file theo tuỳ chọn của người dùng (hộp thoại xuất).

Nhận DỮ LIỆU ĐÃ SỬA từ trình duyệt chứ không đọc lại kết quả quét: người dùng
sửa ô nào trên màn hình thì file tải về phải đúng y như vậy. Đây là điểm dễ sai
nhất — xuất từ dữ liệu gốc là công soát của họ đổ sông đổ bể.

Tuỳ chọn:
- format:    xlsx | csv
- scope:     danh sách sheet_id cần xuất (frontend đã lọc sẵn)
- structure: per_sheet (mỗi mặt sổ một tab) | merged (gộp một bảng)
- totals:    thêm dòng 合計 ở cuối
"""

from pathlib import Path

from openpyxl import Workbook

from export import ARIAL_BOLD, BAND_FILL, BORDER, HEADER_ROW, MONEY_FMT, _write_sheet

from . import export_csv


def _filter_rows(accounts: list[dict], sheet_ids: set[str] | None) -> list[dict]:
    """Giữ lại các dòng thuộc những mặt sổ được chọn, bỏ sổ rỗng."""
    if sheet_ids is None:
        return accounts
    out = []
    for acc in accounts:
        rows = [t for t in acc["transactions"] if t.get("sheet_id") in sheet_ids]
        if rows:
            out.append({**acc, "transactions": rows})
    return out


def _group_per_sheet(
    accounts: list[dict], order: list[str], names: dict[str, str]
) -> list[dict]:
    """Mỗi MẶT SỔ thành một khối riêng, giữ đúng thứ tự người dùng đang thấy."""
    blocks: list[dict] = []
    for sid in order:
        rows = [
            t
            for acc in accounts
            for t in acc["transactions"]
            if t.get("sheet_id") == sid
        ]
        if not rows:
            continue
        src = next(
            (acc for acc in accounts if any(t.get("sheet_id") == sid for t in acc["transactions"])),
            {},
        )
        blocks.append({
            "bank_name": names.get(sid) or src.get("bank_name"),
            "store_no": src.get("store_no"),
            "account_no": src.get("account_no"),
            "holder": src.get("holder"),
            "transactions": rows,
        })
    return blocks


def _merge_all(accounts: list[dict]) -> list[dict]:
    """Nối tiếp mọi giao dịch vào một bảng duy nhất."""
    rows = [t for acc in accounts for t in acc["transactions"]]
    first = accounts[0] if accounts else {}
    return [{
        "bank_name": first.get("bank_name"),
        "store_no": first.get("store_no"),
        "account_no": first.get("account_no"),
        "holder": first.get("holder"),
        "transactions": rows,
    }]


def _append_totals(ws, acc: dict) -> None:
    """Dòng 合計 cuối bảng: cộng 出金/入金 của giao dịch thật (bỏ dòng ghi chú)."""
    txs = [t for t in acc["transactions"] if not t.get("is_note")]
    row = HEADER_ROW + len(acc["transactions"]) + 1

    total_out = sum(t.get("amount_out") or 0 for t in txs)
    total_in = sum(t.get("amount_in") or 0 for t in txs)
    last_bal = next(
        (t["balance"] for t in reversed(acc["transactions"]) if t.get("balance") is not None),
        None,
    )

    headers = [c.value for c in ws[HEADER_ROW]]

    def col_of(name: str) -> int | None:
        for i, h in enumerate(headers, 1):
            if h and str(h).startswith(name):
                return i
        return None

    label_col = col_of("摘要") or 2
    ws.cell(row=row, column=label_col, value="合計").font = ARIAL_BOLD
    for name, value in (
        ("お支払金額", total_out),
        ("お預り金額", total_in),
        ("差引残高", last_bal),
    ):
        c = col_of(name)
        if c and value is not None:
            cell = ws.cell(row=row, column=c, value=value)
            cell.font = ARIAL_BOLD
            cell.number_format = MONEY_FMT
    for i in range(1, len(headers) + 1):
        ws.cell(row=row, column=i).fill = BAND_FILL
        ws.cell(row=row, column=i).border = BORDER


def build(
    accounts: list[dict],
    *,
    fmt: str = "xlsx",
    sheet_ids: list[str] | None = None,
    sheet_order: list[str] | None = None,
    sheet_names: dict[str, str] | None = None,
    structure: str = "per_sheet",
    totals: bool = False,
    book_name: str = "",
    out_path: Path | None = None,
) -> bytes | str:
    """Dựng file. Trả bytes cho CSV, đường dẫn file cho XLSX."""
    picked = _filter_rows(accounts, set(sheet_ids) if sheet_ids else None)

    if structure == "per_sheet" and sheet_order:
        blocks = _group_per_sheet(picked, sheet_order, sheet_names or {})
    elif structure == "merged":
        blocks = _merge_all(picked)
    else:
        blocks = picked

    if not blocks:
        raise ValueError("Không có dòng nào để xuất — hãy chọn mặt sổ đã quét.")

    if fmt == "csv":
        return export_csv.accounts_to_csv(blocks, book_name, totals=totals)

    wb = Workbook()
    used: set[str] = set()
    for i, acc in enumerate(blocks):
        title = (acc.get("bank_name") or f"口座{i + 1}")[:28]
        for ch in '\\/*?:[]':
            title = title.replace(ch, "")
        title = title.strip() or f"口座{i + 1}"
        if title in used:
            title = f"{title[:25]}({i + 1})"
        used.add(title)

        sheet = wb.active if i == 0 else wb.create_sheet()
        sheet.title = title
        _write_sheet(sheet, acc, book_name)
        if totals:
            _append_totals(sheet, acc)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return str(out)
